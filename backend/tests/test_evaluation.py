"""Focused evaluation acceptance cases. Execution requires project-owner approval."""

import asyncio
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from openpyxl import Workbook, load_workbook

from app.schemas.evaluation import ExpectedAnswer, JudgeVerdict, ReviewInput
from app.services.config_service import ConfigError
from app.services.evaluation import service
from app.services.evaluation.metrics import labelled_metrics
from app.services.evaluation.rules import evaluate_rules, semantic_signature
from app.services.evaluation.workbook import make_workbook, parse_workbook


def action(index, route="training_business", intent="start", slots=None, dependencies=None):
    return {
        "action_id": f"action_{index}",
        "order": index,
        "route": route,
        "intent": {"name": intent, "zh": intent, "confidence": 0.9} if intent else None,
        "slots": slots or [],
        "depends_on": dependencies or [],
    }


def result(actions=None, composition="single"):
    return {
        "schema_version": 2,
        "request_id": "request",
        "text": "问题",
        "actions": actions or [action(1)],
        "composition": composition,
        "is_compound": composition != "single",
        "risk": None,
        "need_clarification": False,
        "clarification": None,
    }


def case(**values):
    return {
        "case_id": "a",
        "question": "问题",
        "standard_answer": "标准答案",
        "tags": [],
        "critical": False,
        "context": {},
        **values,
    }


def measured(status, *, index=0, critical=False):
    return {
        **case(critical=critical),
        "id": index + 1,
        "repeat_index": index,
        "status": status,
        "actual": result() if status != "error" else None,
        "latency_ms": 100,
        "issue_tags": [],
        "verdict": {},
    }


def test_parallel_permutation_is_not_a_false_failure():
    expected = {
        "composition": "parallel",
        "actions": [
            {"route": "training_business", "intent": "start"},
            {"route": "dialogue_device_control", "intent": "connect"},
        ],
    }
    actual = result([action(1, "dialogue_device_control", "connect"), action(2)], "parallel")
    assert evaluate_rules(expected, actual)["passed"]


def test_one_actual_action_cannot_satisfy_two_expected_actions():
    expected = {"actions": [{"route": "training_business"}, {"route": "training_business"}]}
    judged = evaluate_rules(expected, result())
    assert not judged["passed"]
    assert any(issue["tag"] == "action_missing" for issue in judged["issues"])


def test_slots_in_wrong_actions_are_not_accepted():
    expected = {
        "composition": "sequential",
        "actions": [
            {"route": "training_business", "slots": [{"name": "exercise", "value": "跑步"}]},
            {
                "route": "training_business",
                "slots": [{"name": "exercise", "value": "深蹲"}],
                "depends_on": [1],
            },
        ],
    }
    actual = result(
        [
            action(1, slots=[{"name": "exercise", "value": "深蹲"}]),
            action(2, slots=[{"name": "exercise", "value": "跑步"}], dependencies=["action_1"]),
        ],
        "sequential",
    )
    assert [issue["tag"] for issue in evaluate_rules(expected, actual)["issues"]] == [
        "slot",
        "slot",
    ]


def test_dependency_checks_use_action_mapping_not_expected_ids():
    expected = {
        "actions": [
            {"route": "training_business"},
            {"route": "training_business", "depends_on": [1]},
        ]
    }
    actual = result([action(1), action(2)], "sequential")
    assert any(issue["tag"] == "dependency" for issue in evaluate_rules(expected, actual)["issues"])


def test_unlabelled_is_distinct_from_explicit_null():
    no_checks = ExpectedAnswer().model_dump(exclude_unset=True)
    no_intent = ExpectedAnswer(actions=[{"route": "training_business", "intent": None}]).model_dump(
        exclude_unset=True
    )
    assert evaluate_rules(no_checks, result())["passed"]
    assert not evaluate_rules(no_intent, result())["passed"]


@pytest.mark.parametrize("expected,actual", [(True, False), (False, True)])
def test_clarification_is_checked_in_both_directions(expected, actual):
    output = {**result(), "need_clarification": actual}
    assert not evaluate_rules({"need_clarification": expected}, output)["passed"]


def test_stability_ignores_confidence_but_not_slot_values():
    first = result([action(1, slots=[{"name": "exercise", "value": "跑步"}])])
    second = result([action(1, slots=[{"name": "exercise", "value": "跑步"}])])
    second["actions"][0]["intent"]["confidence"] = 0.2
    assert semantic_signature(first) == semantic_signature(second)
    second["actions"][0]["slots"][0]["value"] = "深蹲"
    assert semantic_signature(first) != semantic_signature(second)


def test_errors_remain_in_pass_rate_denominator():
    run = SimpleNamespace(
        status="completed", repeat_count=1, evaluation_config_json={"pass_threshold": 60}
    )
    items = [
        dict(measured(status), case_id=str(index))
        for index, status in enumerate(["pass", "fail", "error"])
    ]
    stats = service.summary(run, items)
    assert stats["total"] == 3 and stats["pass_rate"] == 33.33
    assert stats["final"] and stats["acceptance"] == "failed"


def test_pending_repeat_prevents_final_acceptance_without_reweighting_primary():
    run = SimpleNamespace(
        status="completed", repeat_count=3, evaluation_config_json={"pass_threshold": 100}
    )
    stats = service.summary(
        run, [measured("pass"), measured("fail", index=1), measured("review", index=2)]
    )
    assert stats["pass_rate"] == 100
    assert not stats["final"] and stats["pending_executions"] == 1


def test_critical_repeat_failure_blocks_acceptance():
    run = SimpleNamespace(
        status="completed", repeat_count=3, evaluation_config_json={"pass_threshold": 100}
    )
    stats = service.summary(
        run,
        [
            measured("pass", critical=True),
            measured("fail", index=1, critical=True),
            measured("pass", index=2, critical=True),
        ],
    )
    assert stats["pass_rate"] == 100 and stats["acceptance"] == "failed"


def test_extra_actions_and_execution_errors_affect_route_metrics():
    first = {
        **case(expected={"actions": [{"route": "training_business"}]}),
        "actual": result([action(1), action(2, "safety")], "parallel"),
    }
    second = {**case(expected={"actions": [{"route": "safety"}]}), "actual": None}
    metrics = labelled_metrics([first, second])
    safety = next(entry for entry in metrics["per_route"] if entry["route"] == "safety")
    assert safety["support"] == safety["predicted"] == 1
    assert safety["precision"] == safety["recall"] == 0


def test_two_column_workbook_is_sufficient():
    book = Workbook()
    book.active.title = "用例"
    book.active.append(["问题", "标准答案"])
    book.active.append(["你好", "识别为社交问候，不追问"])
    output = BytesIO()
    book.save(output)
    imported = parse_workbook(output.getvalue(), "简单模板", 500)
    assert len(imported.cases) == 1
    assert imported.cases[0].standard_answer == "识别为社交问候，不追问"


def test_structured_workbook_roundtrip_and_formula_safe_export():
    source = case(
        question="=1+1",
        context={"device": "watch"},
        expected={
            "risk": None,
            "composition": "sequential",
            "actions": [
                {
                    "route": "dialogue_device_control",
                    "intent": "connect",
                    "slots": [],
                    "depends_on": [],
                },
                {
                    "route": "training_business",
                    "slots": [{"name": "exercise", "value": "跑步"}],
                    "depends_on": [1],
                },
            ],
        },
    )
    workbook = make_workbook([source])
    parsed = parse_workbook(workbook, "复合用例", 500).cases[0].model_dump(exclude_unset=True)
    assert parsed == source
    opened = load_workbook(BytesIO(workbook))
    assert opened["用例"]["B2"].data_type == "s"
    opened.close()


def test_import_rejects_formulas_and_reports_row():
    book = Workbook()
    book.active.title = "用例"
    book.active.append(["问题", "标准答案"])
    book.active.append(["=1+1", "答案"])
    output = BytesIO()
    book.save(output)
    with pytest.raises(ConfigError) as captured:
        parse_workbook(output.getvalue(), "公式", 500)
    assert captured.value.issues[0]["field"] == "用例!2"


def test_missing_gold_is_not_silently_imported():
    book = Workbook()
    book.active.title = "用例"
    book.active.append(["问题", "标准答案"])
    book.active.append(["问题", ""])
    output = BytesIO()
    book.save(output)
    with pytest.raises(ConfigError) as captured:
        parse_workbook(output.getvalue(), "缺失标准", 500)
    assert captured.value.issues[0]["field"].startswith("用例!2")


def test_judge_unknown_verdict_cannot_be_accepted():
    with pytest.raises(ValueError):
        JudgeVerdict(verdict="probably_pass", reason="不确定", tags=[])


def test_human_cannot_override_hard_rules(monkeypatch):
    monkeypatch.setattr(service, "require_run", AsyncMock())
    row = SimpleNamespace(
        review_revision=0,
        auto_status="fail",
        verdict_json={"rules": {"issues": [{"tag": "slot"}]}},
        review_json=[],
    )
    session = SimpleNamespace(scalar=AsyncMock(return_value=row))
    with pytest.raises(ConfigError, match="精确规则"):
        asyncio.run(
            service.review_result(
                session,
                1,
                1,
                ReviewInput(decision="pass", reason="理由", reviewer="甲", expected_revision=0),
            )
        )


def test_review_has_optimistic_concurrency(monkeypatch):
    monkeypatch.setattr(service, "require_run", AsyncMock())
    session = SimpleNamespace(scalar=AsyncMock(return_value=SimpleNamespace(review_revision=2)))
    with pytest.raises(ConfigError, match="其他人复核"):
        asyncio.run(
            service.review_result(
                session,
                1,
                1,
                ReviewInput(decision="fail", reason="理由", reviewer="乙", expected_revision=1),
            )
        )
