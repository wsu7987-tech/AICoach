"""Bounded XLSX import/export; the only required product columns are Q and gold."""

import json
from io import BytesIO
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import ValidationError

from app.schemas.evaluation import DatasetCreate
from app.services.config_service import ConfigError

HEADERS = [
    "用例编号",
    "问题",
    "标准答案",
    "场景标签",
    "关键用例",
    "期望组合方式",
    "期望追问",
    "期望风险类型",
    "期望风险等级",
    "上下文（可选）",
]
ACTION_HEADERS = ["用例编号", "动作编号", "期望Route", "期望Intent", "依赖动作", "检查参数"]
SLOT_HEADERS = ["用例编号", "动作编号", "参数名", "参数值"]
COMPOSITIONS = {"单动作": "single", "并行": "parallel", "顺序": "sequential", "备选": "alternative"}
BOOLS = {"是": True, "否": False, "true": True, "false": False}


def _error(message: str, issues: list[dict] | None = None):
    raise ConfigError("EVALUATION_IMPORT_INVALID", message, 422, issues=issues or [])


def _safe_cell(value):
    # Explicit string cells prevent imported/model-produced text becoming formulas.
    return "" if value is None else ILLEGAL_CHARACTERS_RE.sub("", str(value))


def _sheet(book, name: str, headers: list[str], rows: list[list]):
    sheet = book.create_sheet(name)
    sheet.append(headers)
    for row in rows:
        sheet.append([_safe_cell(value) for value in row])
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 25
    if name == "用例":
        sheet.column_dimensions["B"].width = 50
        sheet.column_dimensions["C"].width = 70
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    return sheet


def make_workbook(
    cases: list[dict] | None = None, results: list[dict] | None = None, report: dict | None = None
) -> bytes:
    book = Workbook()
    book.remove(book.active)
    rows, actions, slots = [], [], []
    for case in cases or []:
        expected = case.get("expected") or {}
        risk = expected.get("risk")
        rows.append(
            [
                case["case_id"],
                case["question"],
                case["standard_answer"],
                "，".join(case["tags"]),
                "是" if case.get("critical") else "否",
                expected.get("composition", ""),
                "是"
                if expected.get("need_clarification") is True
                else "否"
                if expected.get("need_clarification") is False
                else "",
                risk["type"] if risk else "无" if "risk" in expected else "",
                risk["level"] if risk else "",
                json.dumps(case["context"], ensure_ascii=False) if case.get("context") else "",
            ]
        )
        for index, action in enumerate(expected.get("actions") or [], 1):
            actions.append(
                [
                    case["case_id"],
                    index,
                    action["route"],
                    action.get("intent") or "无" if "intent" in action else "",
                    ",".join(map(str, action["depends_on"])) or "无"
                    if action.get("depends_on") is not None
                    else "",
                    "是" if action.get("slots") is not None else "",
                ]
            )
            for slot in action.get("slots") or []:
                slots.append([case["case_id"], index, slot["name"], slot["value"]])
    _sheet(book, "用例", HEADERS, rows)
    _sheet(book, "动作（可选）", ACTION_HEADERS, actions)
    _sheet(book, "参数（可选）", SLOT_HEADERS, slots)
    _sheet(
        book,
        "填写说明",
        ["项目", "说明"],
        [
            [
                "必填",
                "用例表只需填写问题、标准答案；用例编号留空会按行生成。不要在用例表保留示例行。",
            ],
            [
                "标准答案",
                "用业务语言说明应该识别的动作、参数、先后关系、风险及是否追问。无需 JSON。",
            ],
            [
                "评审",
                "选评审模型自动判断标准答案；不选则由人工复核。可选结构化规则错误不能被模型或人工覆盖。",
            ],
            ["可选检查", "辅助字段留空表示不检查；无表示明确要求为空。期望追问填写是/否。"],
            [
                "动作",
                "编号从 1 连续递增；Route、Intent 填配置中的英文名称或中文名称/别名。"
                "依赖动作填写前置编号，如 1,2。",
            ],
            [
                "参数",
                "同一动作可填写多行参数；填写参数或检查参数=是后，精确比较全部参数（含多余参数）。枚举值填写标准值。",
            ],
            ["组合方式", "single/单动作、parallel/并行、sequential/顺序、alternative/备选。"],
            [
                "风险",
                "类型填英文风险类型，等级填 low/medium/high；明确无风险时类型填无，等级留空。",
            ],
            ["标签", "多个标签用逗号分隔。关键用例填写是/否，关键用例未通过会阻止验收达标。"],
            ["上下文", "通常留空。需要传入 context 的高级用例填写 JSON 对象；导出会原样保留。"],
            ["示例问题", "帮我连接手表，然后开始跑步训练"],
            [
                "示例标准答案",
                "两个动作：连接手表、开始跑步训练；先连接再训练，第二个动作依赖第一个。",
            ],
        ],
    )
    if results is not None:
        chunks = []

        def result_cell(item, field):
            value = json.dumps(item[field], ensure_ascii=False)
            if len(value) <= 30000:
                return value
            for part, start in enumerate(range(0, len(value), 30000), 1):
                chunks.append(
                    [
                        item["case_id"],
                        item["repeat_index"] + 1,
                        field,
                        part,
                        value[start : start + 30000],
                    ]
                )
            return "内容超过 Excel 单元格长度限制，完整内容见“完整结果分段”工作表"

        _sheet(
            book,
            "评测结果",
            [
                "用例编号",
                "问题",
                "标准答案",
                "场景标签",
                "重复序号",
                "最终判定",
                "自动判定",
                "问题标签",
                "判定原因",
                "实际结果",
                "request_id",
                "trace_id",
                "Trace状态",
                "耗时ms",
                "人工复核记录",
            ],
            [
                [
                    item["case_id"],
                    item["question"],
                    item["standard_answer"],
                    "，".join(item["tags"]),
                    item["repeat_index"] + 1,
                    item["status"],
                    item["auto_status"],
                    "，".join(item["issue_tags"]),
                    item["reason"],
                    result_cell(item, "actual"),
                    item["request_id"],
                    item["trace_id"],
                    item["trace_state"],
                    item["latency_ms"],
                    result_cell(item, "reviews"),
                ]
                for item in results
            ],
        )
        if chunks:
            _sheet(
                book,
                "完整结果分段",
                ["用例编号", "重复序号", "字段", "分段序号", "内容（按序拼接）"],
                chunks,
            )
    if report is not None:
        _sheet(book, "任务概览", ["项目", "值"], [[key, value] for key, value in report.items()])
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def parse_workbook(content: bytes, name: str, max_cases: int) -> DatasetCreate:
    try:
        with ZipFile(BytesIO(content)) as archive:
            if (
                sum(info.file_size for info in archive.infolist()) > 64 * 1024 * 1024
                or len(archive.infolist()) > 200
            ):
                _error("工作簿解压后过大，请缩小文件")
        book = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except ConfigError:
        raise
    except Exception as exc:
        raise ConfigError(
            "EVALUATION_IMPORT_INVALID", "无法读取 Excel，请上传有效的 .xlsx 文件", 422
        ) from exc
    try:
        if "用例" not in book.sheetnames:
            _error("缺少“用例”工作表，请下载模板")
        issues: list[dict] = []

        def rows(sheet_name: str, allowed: list[str], required: list[str], limit: int):
            if sheet_name not in book.sheetnames:
                return []
            sheet = book[sheet_name]
            if (sheet.max_row or 0) > limit + 1 or (sheet.max_column or 0) > 20:
                _error(f"{sheet_name} 工作表行数或列数超过限制")
            data = sheet.iter_rows()
            header = next(data, ())
            names = [str(cell.value).strip() if cell.value is not None else "" for cell in header]
            if any(column not in names for column in required):
                _error(f"{sheet_name} 缺少必填列：{'、'.join(required)}")
            if len([value for value in names if value]) != len(
                set(value for value in names if value)
            ):
                _error(f"{sheet_name} 存在重复列名")
            if any(value and value not in allowed for value in names):
                _error(f"{sheet_name} 存在未知列，请使用当前模板")
            output = []
            for number, row in enumerate(data, 2):
                if number > limit + 1:
                    _error(f"{sheet_name} 超过行数限制")
                if not any(cell.value is not None for cell in row):
                    continue
                if any(cell.data_type == "f" for cell in row):
                    issues.append(
                        {"field": f"{sheet_name}!{number}", "message": "不接受公式，请粘贴为文本值"}
                    )
                values = {
                    names[index]: str(cell.value).strip() if cell.value is not None else ""
                    for index, cell in enumerate(row)
                    if index < len(names) and names[index]
                }
                output.append((number, values))
            return output

        def boolean(value, field):
            if value.lower() not in BOOLS:
                issues.append({"field": field, "message": "请填写是或否"})
                return False
            return BOOLS[value.lower()]

        cases: dict[str, dict] = {}
        case_lines = []
        for number, row in rows("用例", HEADERS, ["问题", "标准答案"], max_cases):
            key = row.get("用例编号") or f"case_{number - 1}"
            if key in cases:
                issues.append({"field": f"用例!{number}", "message": f"用例编号 {key} 重复"})
                continue
            expected = {}
            if row.get("期望组合方式"):
                expected["composition"] = COMPOSITIONS.get(row["期望组合方式"], row["期望组合方式"])
            if row.get("期望追问"):
                expected["need_clarification"] = boolean(row["期望追问"], f"用例!{number}.期望追问")
            if row.get("期望风险类型"):
                expected["risk"] = (
                    None
                    if row["期望风险类型"] == "无"
                    else {"type": row["期望风险类型"], "level": row.get("期望风险等级", "")}
                )
            elif row.get("期望风险等级"):
                issues.append(
                    {
                        "field": f"用例!{number}.期望风险类型",
                        "message": "填写风险等级时必须填写风险类型",
                    }
                )
            context = {}
            if row.get("上下文（可选）"):
                try:
                    context = json.loads(row["上下文（可选）"])
                    if not isinstance(context, dict):
                        raise ValueError()
                except (ValueError, TypeError):
                    issues.append(
                        {"field": f"用例!{number}.上下文", "message": "上下文必须是 JSON 对象"}
                    )
            case_lines.append(number)
            cases[key] = {
                "case_id": key,
                "question": row["问题"],
                "standard_answer": row["标准答案"],
                "tags": [
                    value.strip()
                    for value in row.get("场景标签", "").replace("，", ",").split(",")
                    if value.strip()
                ],
                "critical": boolean(row["关键用例"], f"用例!{number}.关键用例")
                if row.get("关键用例")
                else False,
                "context": context,
                "expected": expected,
            }

        action_rows: dict[tuple[str, int], dict] = {}
        for number, row in rows("动作（可选）", ACTION_HEADERS, ACTION_HEADERS[:3], max_cases * 12):
            try:
                key, index = row["用例编号"], int(row["动作编号"])
                if key not in cases or not 1 <= index <= 12 or (key, index) in action_rows:
                    raise ValueError()
                action = {"route": row["期望Route"]}
                if row.get("期望Intent"):
                    action["intent"] = None if row["期望Intent"] == "无" else row["期望Intent"]
                if row.get("依赖动作"):
                    action["depends_on"] = (
                        []
                        if row["依赖动作"] == "无"
                        else [
                            int(value.strip())
                            for value in row["依赖动作"].replace("，", ",").split(",")
                        ]
                    )
                if row.get("检查参数") and boolean(row["检查参数"], f"动作!{number}.检查参数"):
                    action["slots"] = []
                action_rows[key, index] = action
            except (ValueError, KeyError):
                issues.append(
                    {
                        "field": f"动作!{number}",
                        "message": "用例不存在、动作编号重复/无效或依赖编号格式错误",
                    }
                )

        for number, row in rows("参数（可选）", SLOT_HEADERS, SLOT_HEADERS, max_cases * 50):
            try:
                action = action_rows[row["用例编号"], int(row["动作编号"])]
                action.setdefault("slots", []).append(
                    {"name": row["参数名"], "value": row["参数值"]}
                )
            except (ValueError, KeyError):
                issues.append({"field": f"参数!{number}", "message": "引用的用例或动作不存在"})
        for key, case in cases.items():
            indexes = sorted(index for case_id, index in action_rows if case_id == key)
            if indexes:
                if indexes != list(range(1, len(indexes) + 1)):
                    issues.append({"field": f"动作.{key}", "message": "动作编号必须从 1 连续递增"})
                case["expected"]["actions"] = [action_rows[key, index] for index in indexes]
        if issues:
            _error("导入校验失败，请修正标记的行", issues[:100])
        try:
            return DatasetCreate(name=name, cases=list(cases.values()))
        except ValidationError as exc:
            _error(
                "问题、标准答案或辅助字段填写不完整",
                [
                    {
                        "field": (
                            f"用例!{case_lines[error['loc'][1]]}."
                            + ".".join(map(str, error["loc"][2:]))
                            if len(error["loc"]) > 1
                            and error["loc"][0] == "cases"
                            and isinstance(error["loc"][1], int)
                            else ".".join(map(str, error["loc"]))
                        ),
                        "message": error["msg"],
                    }
                    for error in exc.errors(include_input=False, include_url=False)
                ][:100],
            )
    finally:
        book.close()
